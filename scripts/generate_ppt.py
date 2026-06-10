"""
Auto_Upt 路演 PPT 生成脚本
生成 13 页可编辑的 .pptx 文件，可直接用 PowerPoint 打开编辑。
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── 主题色彩 ──
PRIMARY = RGBColor(0x1A, 0x56, 0xDB)       # 深蓝主色
PRIMARY_LIGHT = RGBColor(0x3B, 0x82, 0xF6)  # 亮蓝
ACCENT = RGBColor(0xF5, 0x9E, 0x0B)         # 琥珀强调
ACCENT_GREEN = RGBColor(0x10, 0xB9, 0x81)    # 绿色（已完成）
ACCENT_ORANGE = RGBColor(0xF5, 0x9E, 0x0B)   # 橙色（进行中）
ACCENT_GRAY = RGBColor(0x9C, 0xA3, 0xAF)      # 灰色（未完成）
DARK = RGBColor(0x1F, 0x29, 0x37)             # 深色文字
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
CARD_BG = RGBColor(0xEF, 0xF6, 0xFF)
BORDER = RGBColor(0xE5, 0xE7, 0xEB)

OUTPUT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_PATH = os.path.join(OUTPUT_DIR, "docs", "Auto_Upt_路演PPT.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 宽屏
prs.slide_height = Inches(7.5)

# ── 工具函数 ──

def add_blank_slide():
    """添加空白幻灯片"""
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def add_bg(slide, color=WHITE):
    """设置幻灯片背景色"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill_color=None, border_color=None):
    """添加矩形"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=18, color=DARK,
                 bold=False, alignment=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    """添加文本框"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_multiline_text(slide, left, top, width, height, lines, font_size=16, color=DARK,
                       bold_first=False, line_spacing=1.5, font_name="Microsoft YaHei",
                       alignment=PP_ALIGN.LEFT):
    """添加多行文本框，每行一个段落"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = font_name
        p.alignment = alignment
        p.space_after = Pt(font_size * (line_spacing - 1))
        if bold_first and i == 0:
            p.font.bold = True
    return txBox


def add_icon_circle(slide, left, top, size, color_fill, text, font_size=24):
    """添加圆形图标 + 文字"""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color_fill
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = "Microsoft YaHei"
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(0)
    return shape


def add_section_header(slide, number, title, subtitle=""):
    """添加左侧装饰条 + 编号 + 标题"""
    # 左侧色条
    add_rect(slide, Inches(0), Inches(0), Inches(0.15), Inches(7.5), fill_color=PRIMARY)
    # 顶部装饰线
    add_rect(slide, Inches(0.15), Inches(0), Inches(13.18), Inches(0.04), fill_color=PRIMARY_LIGHT)
    # 底部装饰线
    add_rect(slide, Inches(0.15), Inches(7.46), Inches(13.18), Inches(0.04), fill_color=BORDER)
    # 编号圆
    add_icon_circle(slide, Inches(0.6), Inches(0.4), Inches(0.55), PRIMARY, str(number), font_size=22)
    # 标题
    add_text_box(slide, Inches(1.35), Inches(0.42), Inches(10), Inches(0.7), title,
                 font_size=28, color=DARK, bold=True)
    if subtitle:
        add_text_box(slide, Inches(1.35), Inches(1.0), Inches(10), Inches(0.4), subtitle,
                     font_size=14, color=ACCENT_GRAY)


def add_card(slide, left, top, width, height, title, body_lines, accent_color=PRIMARY):
    """添加卡片组件"""
    # 顶部色条
    add_rect(slide, left, top, width, Inches(0.06), fill_color=accent_color)
    # 卡片主体
    add_rect(slide, left, top + Inches(0.06), width, height - Inches(0.06),
             fill_color=CARD_BG, border_color=BORDER)
    # 标题
    add_text_box(slide, left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), Inches(0.4),
                 title, font_size=16, color=accent_color, bold=True)
    # 正文
    if body_lines:
        add_multiline_text(slide, left + Inches(0.2), top + Inches(0.65), width - Inches(0.4),
                           height - Inches(0.8), body_lines, font_size=13, color=DARK, line_spacing=1.4)


def add_step_number(slide, left, top, size, number):
    """添加步骤编号圆"""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PRIMARY
    shape.line.fill.background()
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = str(number)
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = "Microsoft YaHei"
    p.alignment = PP_ALIGN.CENTER
    return shape


def add_page_number(slide, num):
    """右下角页码"""
    add_text_box(slide, Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.35),
                 str(num), font_size=10, color=ACCENT_GRAY, alignment=PP_ALIGN.RIGHT)


# ═══════════════════════════════════════════
#  幻灯片 1：封面
# ═══════════════════════════════════════════
slide1 = add_blank_slide()
add_bg(slide1, DARK)
# 顶部装饰线
add_rect(slide1, Inches(0), Inches(0), Inches(13.333), Inches(0.08), fill_color=PRIMARY_LIGHT)
# 底部装饰线
add_rect(slide1, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08), fill_color=ACCENT)
# 主标题
add_text_box(slide1, Inches(1.5), Inches(1.5), Inches(10.3), Inches(1.2),
             "Auto_Upt", font_size=56, color=WHITE, bold=True)
# 副标题
add_text_box(slide1, Inches(1.5), Inches(2.8), Inches(10.3), Inches(0.8),
             "多平台创作内容自动发布工具", font_size=28, color=PRIMARY_LIGHT)
# 一句话定位
add_text_box(slide1, Inches(1.5), Inches(3.7), Inches(10.3), Inches(0.7),
             "一次创作，四处发布，AI 帮你适配每个平台", font_size=20, color=ACCENT_GRAY)
# 中间分割线
add_rect(slide1, Inches(1.5), Inches(4.7), Inches(3), Inches(0.03), fill_color=PRIMARY_LIGHT)
# 日期 & 团队
add_text_box(slide1, Inches(1.5), Inches(5.0), Inches(5), Inches(0.4),
             "2026 年 6 月  |  路演演示", font_size=16, color=ACCENT_GRAY)
add_page_number(slide1, 1)


# ═══════════════════════════════════════════
#  幻灯片 2：创业者真实痛点
# ═══════════════════════════════════════════
slide2 = add_blank_slide()
add_bg(slide2, WHITE)
add_section_header(slide2, 2, "创作者的真实痛点", "跨平台发布的重复劳动到底有多重？")

# 四个痛点卡片
pain_points = [
    ("📏 标题长度限制不同", "公众号 ≤64 字\nB站 ≤80 字\n知乎 无硬限制\n小红书 ≤20 字"),
    ("📝 正文风格不同", "公众号：深度长文\n知乎：专业问答\nB站：视频简介向\n小红书：种草笔记风"),
    ("🖼️ 图片比例不兼容", "公众号：16:9 封面\nB站：16:9 + 竖版\n小红书：1:1 / 3:4\n知乎：无限制"),
    ("🏷️ 标签规则各异", "公众号：无标签\nB站：分区 + 标签\n小红书：话题标签\n知乎：话题绑定"),
]

card_w = Inches(2.8)
card_h = Inches(4.5)
start_x = Inches(0.6)
gap = Inches(0.25)
start_y = Inches(2.0)

for i, (title, body) in enumerate(pain_points):
    x = start_x + (card_w + gap) * i
    # 顶部色条
    add_rect(slide2, x, start_y, card_w, Inches(0.06), fill_color=PRIMARY_LIGHT)
    # 卡片背景
    add_rect(slide2, x, start_y + Inches(0.06), card_w, card_h - Inches(0.06),
             fill_color=LIGHT_BG, border_color=BORDER)
    add_text_box(slide2, x + Inches(0.2), start_y + Inches(0.25), card_w - Inches(0.4), Inches(0.5),
                 title, font_size=15, color=DARK, bold=True)
    add_multiline_text(slide2, x + Inches(0.2), start_y + Inches(0.85), card_w - Inches(0.4),
                       card_h - Inches(1.0), body.split("\n"), font_size=12, color=DARK, line_spacing=1.6)

# 底部总结
add_text_box(slide2, Inches(0.6), Inches(6.8), Inches(12), Inches(0.4),
             "💡 手动逐个平台改写 ≈ 1~2 小时 / 篇，跨平台创作者每周花数十小时在重复发布上",
             font_size=15, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)
add_page_number(slide2, 2)


# ═══════════════════════════════════════════
#  幻灯片 3：我们的答案
# ═══════════════════════════════════════════
slide3 = add_blank_slide()
add_bg(slide3, WHITE)
add_section_header(slide3, 3, "我们的答案", "内容中台 + 平台适配器 + AI Agent 发布助理")

# 三大支柱
pillars = [
    ("🏗️", "内容中台", "用户只维护一份统一内容\nMarkdown 输入 → 内容 IR\n统一管理素材、标签、发布字段"),
    ("🔌", "平台适配器", "公众号 / B站 / 知乎 / 小红书\n每个平台独立渲染 & 校验\n新增平台只需 3 个文件"),
    ("🤖", "AI Agent", "内容分析 → 风格改写 → 格式检查\n合规审查 → 发布建议\n不绕过人工确认，你是最终决策者"),
]

for i, (icon, title, body) in enumerate(pillars):
    x = Inches(1.0) + Inches(4.0) * i
    # 图标圆
    add_icon_circle(slide3, x + Inches(1.3), Inches(2.0), Inches(1.2), PRIMARY, icon, font_size=40)
    add_text_box(slide3, x, Inches(3.4), Inches(3.8), Inches(0.5), title,
                 font_size=22, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
    add_multiline_text(slide3, x + Inches(0.3), Inches(4.0), Inches(3.2), Inches(2.0),
                       body.split("\n"), font_size=13, color=DARK, line_spacing=1.6,
                       alignment=PP_ALIGN.CENTER)

# 底部一句话
add_text_box(slide3, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.5),
             "核心理念：你只写一次，系统帮你适配每个平台。但最终的发布按钮，永远在你手里。",
             font_size=16, color=PRIMARY, bold=True, alignment=PP_ALIGN.CENTER)
add_page_number(slide3, 3)


# ═══════════════════════════════════════════
#  幻灯片 4：统一编辑区 → 四平台预览
# ═══════════════════════════════════════════
slide4 = add_blank_slide()
add_bg(slide4, WHITE)
add_section_header(slide4, 4, "统一编辑区 → 一键生成四平台预览",
                   "输入一篇文章，秒出四种平台的差异化渲染")

# 左侧：编辑区示意
add_card(slide4, Inches(0.5), Inches(1.8), Inches(3.5), Inches(4.8),
         "📝 统一编辑区",
         ["标题：AI 内容运营助手如何",
          "       减少重复发布工作",
          "",
          "正文（Markdown）：",
          "## 为什么需要多平台适配",
          "随着内容平台增多...",
          "...",
          "",
          "🏷️ 标签：AI, 内容运营, 自动化",
          "🖼️ 素材：1 封面 + 3 正文图",
          "",
          "✅ 目标平台：",
          "☑ 公众号  ☑ B站",
          "☑ 知乎    ☑ 小红书"],
         accent_color=PRIMARY)

# 右侧：四平台预览
platforms = [
    ("公众号", "长文版", "深度排版，支持富文本元素"),
    ("B站", "视频简介版", "适合视频简介 / 动态发布"),
    ("知乎", "回答/专栏版", "专业问答风格，自动生成摘要"),
    ("小红书", "图文笔记版", "种草笔记风，适配字数限制"),
]
for i, (name, version, desc) in enumerate(platforms):
    col = i % 2
    row = i // 2
    x = Inches(4.5) + Inches(4.2) * col
    y = Inches(1.8) + Inches(2.5) * row
    add_card(slide4, x, y, Inches(3.9), Inches(2.2),
             f"📱 {name} {version}",
             [desc, "→ 自动截断超限标题", "→ 图片按比例适配", "→ 标签映射平台规则"],
             accent_color=PRIMARY_LIGHT)

# 箭头提示
add_text_box(slide4, Inches(3.6), Inches(3.8), Inches(0.8), Inches(0.5),
             "→", font_size=36, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

add_page_number(slide4, 4)


# ═══════════════════════════════════════════
#  幻灯片 5：Agent 智能优化
# ═══════════════════════════════════════════
slide5 = add_blank_slide()
add_bg(slide5, WHITE)
add_section_header(slide5, 5, "Agent 智能优化流水线",
                   "5 个 Agent 各司其职，规则引擎兜底 + LLM 增强")

agents = [
    ("01", "Content Analyst", "分析原文结构、摘要、\n主题、内容类型和素材需求", PRIMARY),
    ("02", "Platform Stylist", "按公众号/知乎/B站/小红书\n风格分别改写标题和正文", PRIMARY_LIGHT),
    ("03", "Format Linter", "检查标题长度、图片比例、\n标签数量、链接限制等", ACCENT_GREEN),
    ("04", "Compliance Reviewer", "AIGC 标注、版权风险、\n夸大宣传、低质营销风险", ACCENT_ORANGE),
    ("05", "Recovery Agent", "发布失败后分析日志和截图\n给出修复建议", ACCENT_GRAY),
]

for i, (num, name, desc, color) in enumerate(agents):
    x = Inches(0.4) + Inches(2.55) * i
    y = Inches(2.0)
    # 编号
    add_icon_circle(slide5, x + Inches(0.8), y, Inches(0.7), color, num, font_size=20)
    # 名称
    add_text_box(slide5, x, y + Inches(0.9), Inches(2.3), Inches(0.5), name,
                 font_size=14, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
    # 描述
    add_text_box(slide5, x + Inches(0.1), y + Inches(1.4), Inches(2.1), Inches(1.2), desc,
                 font_size=11, color=DARK, alignment=PP_ALIGN.CENTER)

# 底部流程箭头
add_text_box(slide5, Inches(0.5), Inches(4.0), Inches(12.3), Inches(0.5),
             "原始输入  ──→  分析  ──→  改写  ──→  校验  ──→  审查  ──→  预览报告",
             font_size=16, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

# 关键提示
add_text_box(slide5, Inches(1.5), Inches(5.2), Inches(10.3), Inches(0.8),
             "⚡ 规则引擎默认可用  |  🔑 配置 OPENAI_API_KEY 后可启用 LLM 增强  |  🛡️ LLM 失败自动回退到规则结果",
             font_size=14, color=DARK, alignment=PP_ALIGN.CENTER)
add_text_box(slide5, Inches(1.5), Inches(6.0), Inches(10.3), Inches(0.5),
             "⚠️ Agent 参与优化，但不绕过人工确认 — 你始终是最终决策者",
             font_size=15, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)
add_page_number(slide5, 5)


# ═══════════════════════════════════════════
#  幻灯片 6：素材库 & 发布确认
# ═══════════════════════════════════════════
slide6 = add_blank_slide()
add_bg(slide6, WHITE)
add_section_header(slide6, 6, "共享素材库 & 发布确认", "从素材管理到发布确认的完整链路")

# 左侧：素材库
add_card(slide6, Inches(0.5), Inches(1.8), Inches(5.5), Inches(4.8),
         "📦 共享素材库",
         ["📁 多级文件夹管理", "🖼️ 图片 / 🎬 视频 / 🔊 音频", "🖼️ 封面图独立管理",
          "🔗 正文素材引用（拖拽插入）", "💾 IndexedDB + localStorage 本地持久化",
          "", "素材全局可用，所有平台共享同一素材库"],
         accent_color=PRIMARY)

# 右侧：发布确认
add_card(slide6, Inches(6.5), Inches(1.8), Inches(6.3), Inches(2.2),
         "✅ 发布确认 — 统一配置",
         ["一个配置，同步所有平台", "标题 / 摘要 / 标签 / 封面 全局生效",
          "适合内容各平台一致的场景"],
         accent_color=PRIMARY_LIGHT)

add_card(slide6, Inches(6.5), Inches(4.3), Inches(6.3), Inches(2.3),
         "✅ 发布确认 — 分平台独立配置",
         ["每个平台独立调整字段", "公众号用完整标题，小红书用短标题",
          "B站加视频链接，知乎加话题绑定", "字段合并规则：用户确认 > Agent 草稿 > 编辑区"],
         accent_color=ACCENT_ORANGE)

add_page_number(slide6, 6)


# ═══════════════════════════════════════════
#  幻灯片 7：任务看板 & 真实发布
# ═══════════════════════════════════════════
slide7 = add_blank_slide()
add_bg(slide7, WHITE)
add_section_header(slide7, 7, "任务看板 & 真实发布追踪", "每条发布任务可追溯、可重试、可查看完整历程")

# 发布流程步骤
steps = [
    ("①", "创建任务", "配置发布字段\n确认提交"),
    ("②", "素材上传", "封面/图片/视频\n上传至目标平台"),
    ("③", "草稿创建", "平台草稿箱\n/ 待审稿件"),
    ("④", "提交发布", "群发 / 提交审核\n/ 扫码确认"),
    ("⑤", "状态追踪", "实时状态查询\n错误日志回溯"),
]
for i, (num, title, desc) in enumerate(steps):
    x = Inches(0.5) + Inches(2.55) * i
    add_icon_circle(slide7, x + Inches(0.85), Inches(2.0), Inches(0.65), PRIMARY, num, font_size=18)
    add_text_box(slide7, x, Inches(2.85), Inches(2.35), Inches(0.4), title,
                 font_size=14, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide7, x + Inches(0.15), Inches(3.25), Inches(2.05), Inches(0.8), desc,
                 font_size=11, color=DARK, alignment=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        add_text_box(slide7, x + Inches(2.1), Inches(2.2), Inches(0.5), Inches(0.5),
                     "→", font_size=24, color=ACCENT_GRAY, bold=True)

# 平台状态表
add_text_box(slide7, Inches(0.5), Inches(4.4), Inches(3), Inches(0.4),
             "各平台真实发布状态：", font_size=15, color=DARK, bold=True)

status_data = [
    ("公众号", "✅ 已接入", "素材上传 → 草稿创建 → 草稿发布 → 状态查询", ACCENT_GREEN),
    ("B站", "✅ 已接入", "视频/封面上传 → 稿件提交 → 状态查询 → 测试删除", ACCENT_GREEN),
    ("小红书", "🔄 部分接入", "myaibot API Adapter 已实现；账号入口和联调待补齐", ACCENT_ORANGE),
    ("知乎", "⏳ 待接入", "当前仅支持预览和模拟发布", ACCENT_GRAY),
]

for i, (platform, status, detail, color) in enumerate(status_data):
    y = Inches(4.9) + Inches(0.55) * i
    add_rect(slide7, Inches(0.5), y, Inches(12.3), Inches(0.48),
             fill_color=LIGHT_BG, border_color=BORDER)
    add_text_box(slide7, Inches(0.7), y + Inches(0.05), Inches(1.8), Inches(0.4),
                 platform, font_size=13, color=DARK, bold=True)
    add_text_box(slide7, Inches(2.6), y + Inches(0.05), Inches(1.8), Inches(0.4),
                 status, font_size=13, color=color, bold=True)
    add_text_box(slide7, Inches(4.5), y + Inches(0.05), Inches(7.5), Inches(0.4),
                 detail, font_size=12, color=DARK)

add_page_number(slide7, 7)


# ═══════════════════════════════════════════
#  幻灯片 8：系统架构总览
# ═══════════════════════════════════════════
slide8 = add_blank_slide()
add_bg(slide8, WHITE)
add_section_header(slide8, 8, "系统架构总览", "内容中台 + 平台适配器 + AI Agent 发布助理")

# 架构层次：从顶部到底部
layers = [
    ("前端工作台", "Vue 3 + Element Plus  |  编辑 / 预览 / 素材库 / 任务看板 / 账号管理", Inches(1.5), PRIMARY),
    ("API 层", "FastAPI  |  RESTful 接口  |  请求校验 → 业务逻辑 → 返回结果", Inches(2.3), PRIMARY_LIGHT),
    ("Agent 编排层", "Content Analyst → Platform Stylist → Format Linter → Compliance Reviewer", Inches(3.1), ACCENT_GREEN),
    ("服务层", "导入服务 / 预览服务 / 发布服务 / 账号服务 / 素材服务", Inches(3.9), ACCENT),
    ("平台适配器", "公众号 Adapter  |  B站 Adapter  |  知乎 Adapter  |  小红书 Adapter", Inches(4.7), PRIMARY_LIGHT),
    ("任务队列", "Celery + Redis  |  异步发布执行  |  重试策略  |  状态回调", Inches(5.5), ACCENT_GREEN),
    ("数据层", "PostgreSQL（账号/内容/发布记录）  |  Redis（队列/缓存）  |  本地存储（素材文件）", Inches(6.3), PRIMARY),
]

for title, desc, y, color in layers:
    add_rect(slide8, Inches(1.0), y, Inches(11.3), Inches(0.65), fill_color=color)
    add_text_box(slide8, Inches(1.3), y + Inches(0.03), Inches(2.5), Inches(0.6),
                 title, font_size=15, color=WHITE, bold=True)
    add_text_box(slide8, Inches(3.9), y + Inches(0.03), Inches(8.0), Inches(0.6),
                 desc, font_size=12, color=WHITE)

add_page_number(slide8, 8)


# ═══════════════════════════════════════════
#  幻灯片 9：平台适配器设计
# ═══════════════════════════════════════════
slide9 = add_blank_slide()
add_bg(slide9, WHITE)
add_section_header(slide9, 9, "平台适配器 — 可扩展架构核心",
                   "新增平台只需 3 个文件：adapter.py / renderer.py / profile.yaml")

# 左侧：Adapter 接口
add_card(slide9, Inches(0.5), Inches(1.8), Inches(6.0), Inches(4.8),
         "🔌 统一 Adapter 接口（5 个职责）",
         ["capabilities   → 声明平台支持的能力（预览/发布/上传）",
          "render         → 内容 IR → 平台草稿（标题/摘要/正文/标签）",
          "validate       → 平台规则校验（字数/图片/标签/合规）",
          "publish        → 真实发布到目标平台（API / 浏览器辅助）",
          "simulate       → 模拟发布（不调用真实 API，用于预览阶段）",
          "",
          "💡 核心服务只依赖接口，不感知平台内部实现"],
         accent_color=PRIMARY)

# 右侧：目录结构
add_card(slide9, Inches(7.0), Inches(1.8), Inches(5.8), Inches(4.8),
         "📁 新增平台只需 3 个文件",
         ["adapters/<新平台>/",
          "  ├── adapter.py     # 实现 5 个接口方法",
          "  ├── renderer.py    # 平台特定渲染逻辑",
          "  └── profile.yaml   # 平台规则配置",
          "",
          "📋 profile.yaml 包含：",
          "• 标题字数限制",
          "• 支持的图片比例",
          "• 标签规则 & 话题规则",
          "• 内容类型（图文/视频/动态）",
          "• 发布模式（API 直发 / 扫码确认）"],
         accent_color=PRIMARY_LIGHT)

add_page_number(slide9, 9)


# ═══════════════════════════════════════════
#  幻灯片 10：技术栈 & 能力矩阵
# ═══════════════════════════════════════════
slide10 = add_blank_slide()
add_bg(slide10, WHITE)
add_section_header(slide10, 10, "技术栈 & 平台能力矩阵", "当前项目技术选型和各平台支持情况一览")

# 技术栈卡片
techs = [
    ("🐍", "Python + FastAPI", "后端 API 服务"),
    ("📋", "Celery + Redis", "异步任务队列"),
    ("🗄️", "PostgreSQL", "持久化存储"),
    ("🎭", "Playwright", "浏览器自动化"),
    ("🖥️", "Vue 3 + Element Plus", "前端工作台"),
    ("🤖", "OpenAI / DeepSeek", "LLM 增强"),
]
for i, (icon, name, role) in enumerate(techs):
    col = i % 3
    row = i // 3
    x = Inches(0.5) + Inches(4.2) * col
    y = Inches(1.8) + Inches(1.3) * row
    add_rect(slide10, x, y, Inches(3.9), Inches(1.1), fill_color=LIGHT_BG, border_color=BORDER)
    add_text_box(slide10, x + Inches(0.15), y + Inches(0.08), Inches(0.5), Inches(0.5),
                 icon, font_size=24)
    add_text_box(slide10, x + Inches(0.7), y + Inches(0.1), Inches(3.0), Inches(0.4),
                 name, font_size=14, color=DARK, bold=True)
    add_text_box(slide10, x + Inches(0.7), y + Inches(0.55), Inches(3.0), Inches(0.35),
                 role, font_size=11, color=ACCENT_GRAY)

# 能力矩阵表
add_text_box(slide10, Inches(0.5), Inches(4.6), Inches(4), Inches(0.4),
             "平台能力矩阵：", font_size=15, color=DARK, bold=True)

matrix_headers = ["平台", "预览", "模拟发布", "真实发布", "Agent 优化"]
matrix_data = [
    ["公众号", "✅", "✅", "✅ 已接入", "✅"],
    ["B站", "✅", "✅", "✅ 已接入", "✅"],
    ["小红书", "✅", "✅", "🔄 部分", "✅"],
    ["知乎", "✅", "✅", "⏳", "✅"],
]

# 表头
for i, h in enumerate(matrix_headers):
    x = Inches(0.5) + Inches(2.5) * i
    add_rect(slide10, x, Inches(5.1), Inches(2.3), Inches(0.45), fill_color=PRIMARY)
    add_text_box(slide10, x, Inches(5.12), Inches(2.3), Inches(0.4), h,
                 font_size=13, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# 数据行
for r, row_data in enumerate(matrix_data):
    y = Inches(5.55) + Inches(0.45) * r
    bg = LIGHT_BG if r % 2 == 0 else WHITE
    for c, val in enumerate(row_data):
        x = Inches(0.5) + Inches(2.5) * c
        add_rect(slide10, x, y, Inches(2.3), Inches(0.45), fill_color=bg, border_color=BORDER)
        fc = DARK
        if c == 0:
            fc = PRIMARY
            fb = True
        else:
            fb = False
        add_text_box(slide10, x, y + Inches(0.04), Inches(2.3), Inches(0.35), val,
                     font_size=12, color=fc, bold=fb, alignment=PP_ALIGN.CENTER)

add_page_number(slide10, 10)


# ═══════════════════════════════════════════
#  幻灯片 11：已完成 vs 进行中
# ═══════════════════════════════════════════
slide11 = add_blank_slide()
add_bg(slide11, WHITE)
add_section_header(slide11, 11, "当前成果：已完成 vs 进行中",
                   "第一阶段 MVP 交付，第二阶段真实发布持续拓展")

# 第一阶段
add_card(slide11, Inches(0.5), Inches(1.8), Inches(6.0), Inches(2.5),
         "✅ 第一阶段 MVP（已完成）",
         ["• 统一内容编辑器（标题/正文/标签/目标平台）",
          "• 共享素材库（图片/视频/音频/多级文件夹）",
          "• 四平台预览（公众号/B站/知乎/小红书）",
          "• Agent 预览编排（规则引擎 + LLM 可选增强）",
          "• 模拟发布任务 & 前端任务看板"],
         accent_color=ACCENT_GREEN)

# 第二阶段
add_card(slide11, Inches(7.0), Inches(1.8), Inches(5.8), Inches(2.5),
         "🔄 第二阶段：真实发布闭环（部分完成）",
         ["• 公众号：素材上传 / 草稿 / 发布 / 状态查询 ✅",
          "• B站：登录凭据 / 稿件提交 / 状态查询 ✅",
          "• 小红书：myaibot API Adapter 已实现 🔄",
          "  ↳ 账号入口 + 真实联调待补齐",
          "• 知乎：暂不进入真实发布 API ⏳"],
         accent_color=ACCENT_ORANGE)

# 待补齐
add_card(slide11, Inches(0.5), Inches(4.6), Inches(12.3), Inches(2.2),
         "📋 当前待补齐项",
         ["• 小红书账号管理入口 & 真实发布联调",
          "• 发布失败时更细粒度的错误恢复建议",
          "• 前端对小红书二维码/状态轮询的完整展示",
          "• 前端 package.json 和自动化测试套件"],
         accent_color=ACCENT)

add_page_number(slide11, 11)


# ═══════════════════════════════════════════
#  幻灯片 12：路线图
# ═══════════════════════════════════════════
slide12 = add_blank_slide()
add_bg(slide12, WHITE)
add_section_header(slide12, 12, "五阶段路线图", "从预览闭环到数据智能反馈的完整演进路径")

phases = [
    ("第一阶段", "预览闭环", "内容编辑 → IR → 多平台预览\n模拟发布 → Agent 编排", ACCENT_GREEN, "✅ 已完成"),
    ("第二阶段", "真实发布", "公众号/B站真实 API 接入\n小红书 API 适配\nCelery 异步发布", ACCENT_ORANGE, "🔄 进行中"),
    ("第三阶段", "浏览器辅助", "知乎 Playwright 自动填充\n小红书备用路线\n验证码人工处理", ACCENT_GRAY, "⏳ 规划中"),
    ("第四阶段", "Agent 工作流", "任务规划/格式检查\n合规检查/失败恢复\n发布报告生成", ACCENT_GRAY, "⏳ 规划中"),
    ("第五阶段", "数据反馈", "阅读/点赞/收藏数据抓取\n数据反哺标题建议\n平台风格持续优化", ACCENT_GRAY, "⏳ 规划中"),
]

for i, (phase, title, desc, color, tag) in enumerate(phases):
    x = Inches(0.4) + Inches(2.55) * i
    y = Inches(2.0)
    # 阶段标识
    add_icon_circle(slide12, x + Inches(0.8), y, Inches(0.65), color, str(i + 1), font_size=18)
    # 连接线
    if i < len(phases) - 1:
        add_rect(slide12, x + Inches(1.5), y + Inches(0.3), Inches(1.0), Inches(0.04),
                 fill_color=BORDER)
    # 标签
    tag_color = ACCENT_GREEN if "已完成" in tag else (ACCENT_ORANGE if "进行中" in tag else ACCENT_GRAY)
    add_text_box(slide12, x + Inches(0.65), y + Inches(0.8), Inches(1.35), Inches(0.3),
                 tag, font_size=9, color=tag_color, bold=True, alignment=PP_ALIGN.CENTER)
    # 标题
    add_text_box(slide12, x, y + Inches(1.2), Inches(2.35), Inches(0.4), title,
                 font_size=15, color=DARK, bold=True, alignment=PP_ALIGN.CENTER)
    # 描述
    add_text_box(slide12, x + Inches(0.1), y + Inches(1.7), Inches(2.15), Inches(1.5), desc,
                 font_size=11, color=DARK, alignment=PP_ALIGN.CENTER)

# 底部时间线
add_rect(slide12, Inches(0.4), Inches(4.9), Inches(12.5), Inches(0.04), fill_color=PRIMARY_LIGHT)
timeline_y = Inches(5.2)
for i, phase in enumerate(phases):
    x = Inches(0.4) + Inches(2.55) * i
    dot_color = ACCENT_GREEN if i == 0 else (ACCENT_ORANGE if i == 1 else ACCENT_GRAY)
    add_icon_circle(slide12, x + Inches(0.95), timeline_y, Inches(0.2), dot_color, "", font_size=1)

add_text_box(slide12, Inches(0.5), Inches(5.8), Inches(12.3), Inches(0.6),
             "当前处于 第一阶段已完成、第二阶段持续推进 的状态。第三至五阶段为未来规划。",
             font_size=14, color=DARK, alignment=PP_ALIGN.CENTER)
add_page_number(slide12, 12)


# ═══════════════════════════════════════════
#  幻灯片 13：总结 & Q&A
# ═══════════════════════════════════════════
slide13 = add_blank_slide()
add_bg(slide13, DARK)
add_rect(slide13, Inches(0), Inches(0), Inches(13.333), Inches(0.08), fill_color=PRIMARY_LIGHT)
add_rect(slide13, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08), fill_color=ACCENT)

add_text_box(slide13, Inches(1.5), Inches(1.0), Inches(10.3), Inches(0.8),
             "总结", font_size=40, color=WHITE, bold=True)

# 三句话回顾
summaries = [
    "1",
    "一次创作 → 多平台自动适配",
    "2",
    "AI 辅助优化，但不绕过人工确认",
    "3",
    "适配器架构，新平台可快速扩展（只需 3 个文件）",
]
for i in range(0, len(summaries), 2):
    row = i // 2
    num_text = summaries[i]
    content_text = summaries[i + 1]
    y = Inches(2.2) + Inches(1.0) * row
    # 编号
    add_icon_circle(slide13, Inches(1.5), y, Inches(0.6), PRIMARY_LIGHT, num_text, font_size=22)
    add_text_box(slide13, Inches(2.4), y + Inches(0.08), Inches(9), Inches(0.5),
                 content_text, font_size=22, color=WHITE)

# 分割线
add_rect(slide13, Inches(1.5), Inches(5.2), Inches(4), Inches(0.03), fill_color=PRIMARY_LIGHT)

# 链接 & 信息
add_text_box(slide13, Inches(1.5), Inches(5.5), Inches(10.3), Inches(0.4),
             "🎬 演示视频", font_size=16, color=PRIMARY_LIGHT, bold=True)
add_text_box(slide13, Inches(1.5), Inches(5.9), Inches(10.3), Inches(0.35),
             "B站：https://www.bilibili.com/video/BV1rbVQ6HEo5/", font_size=13, color=ACCENT_GRAY)
add_text_box(slide13, Inches(1.5), Inches(6.25), Inches(10.3), Inches(0.35),
             "百度网盘：https://pan.baidu.com/s/1oQ5dcYsLzLt2VmYnB0rnCw?pwd=fst3  提取码：fst3",
             font_size=13, color=ACCENT_GRAY)

# Q&A
add_text_box(slide13, Inches(4.5), Inches(6.8), Inches(4.3), Inches(0.5),
             "❓ Q & A", font_size=32, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_page_number(slide13, 13)


# ── 保存 ──
os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
prs.save(OUTPUT_PATH)
print(f"✅ PPT 已生成：{OUTPUT_PATH}")
print(f"   共 {len(prs.slides)} 页幻灯片")
