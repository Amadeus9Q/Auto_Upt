"""文档导入服务 —— 解析 .md / .docx 文件，提取纯文本和媒体信息。"""

from __future__ import annotations

import base64
import io
import logging
import os
import re
from typing import Any
from uuid import uuid4

from fastapi import UploadFile

from backend.app.agents.content_analyst import ContentAnalystAgent
from backend.app.agents.document_extractor import DocumentExtractorAgent
from backend.app.core.config import get_settings
from backend.app.schemas.content import ImportDocumentResponse, ImportedMedia

logger = logging.getLogger(__name__)


class ImportService:
    """文档导入 + LLM 提取服务。"""

    def __init__(self) -> None:
        self._extractor = DocumentExtractorAgent()
        self._analyst = ContentAnalystAgent()

    # ------------------------------------------------------------------
    # 文件解析
    # ------------------------------------------------------------------

    def _parse_upload(self, file: UploadFile) -> tuple[str, list[dict[str, Any]]]:
        """解析上传文件，返回 (纯文本, 媒体资源列表)。"""
        filename = (file.filename or "").lower()

        if filename.endswith(".md") or filename.endswith(".markdown"):
            return self._parse_markdown(file)
        elif filename.endswith(".docx"):
            return self._parse_docx(file)
        elif filename.endswith(".txt"):
            return self._parse_plain_text(file)
        else:
            # 兜底：按纯文本读取
            content = file.file.read().decode("utf-8", errors="replace")
            return content, []

    def _parse_plain_text(self, file: UploadFile) -> tuple[str, list[dict[str, Any]]]:
        """解析纯文本 .txt 文件。"""
        content = file.file.read().decode("utf-8", errors="replace")
        return content, []

    def _parse_markdown(self, file: UploadFile) -> tuple[str, list[dict[str, Any]]]:
        """解析 Markdown 文件。"""
        content = file.file.read().decode("utf-8", errors="replace")
        media: list[dict[str, Any]] = []

        # 提取 ![alt](src) 图片引用，并替换为统一标记
        img_pattern = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
        for i, m in enumerate(img_pattern.finditer(content)):
            media.append({
                "index": i,
                "name": m.group(2).split("/")[-1] or f"image_{i+1}",
                "kind": "image",
                "description": m.group(1) or "",
            })
        content = img_pattern.sub(
            lambda m: f"【图片：{m.group(2).split('/')[-1]}】", content
        )

        return content, media

    def _parse_docx(self, file: UploadFile) -> tuple[str, list[dict[str, Any]]]:
        """解析 .docx 文件，提取纯文本和嵌入图片。"""
        try:
            from docx import Document
        except ImportError:
            logger.warning("python-docx 未安装，按纯文本读取 docx。")
            content = file.file.read().decode("utf-8", errors="replace")
            return content, []

        raw_bytes = file.file.read()
        doc = Document(io.BytesIO(raw_bytes))

        paragraphs: list[str] = []
        media: list[dict[str, Any]] = []
        img_index = 0

        # 提取文档中的图片
        for rel_id, rel in doc.part.rels.items():
            if "image" in rel.reltype:
                try:
                    img_bytes = rel.target_part.blob
                    mime = rel.target_part.content_type or "image/png"
                    ext = mime.split("/")[-1] if "/" in mime else "png"
                    name = os.path.basename(rel.target_part.partname or f"image_{img_index+1}.{ext}")
                    b64 = base64.b64encode(img_bytes).decode("ascii")
                    media.append({
                        "index": img_index,
                        "name": name,
                        "kind": "image",
                        "description": f"文档嵌入图片 {img_index+1}",
                        "mime_type": mime,
                        "base64": b64,
                    })
                    img_index += 1
                except Exception as exc:
                    logger.warning("提取 docx 图片失败: %s", exc)

        # 提取段落文本
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                paragraphs.append("")  # 保留空行分隔
                continue

            # 检测标题样式
            if para.style.name.startswith("Heading"):
                level = para.style.name.split()[-1]
                try:
                    level_num = int(level)
                except ValueError:
                    level_num = 1
                prefix = "#" * min(level_num, 6)
                paragraphs.append(f"{prefix} {text}")
            else:
                paragraphs.append(text)

        raw_text = "\n".join(paragraphs)

        # 处理表格
        for table in doc.tables:
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                raw_text += "\n\n" + "\n".join(rows)

        return raw_text, media

    # ------------------------------------------------------------------
    # 导入入口
    # ------------------------------------------------------------------

    async def import_document(self, file: UploadFile) -> ImportDocumentResponse:
        """导入文档并返回结构化提取结果。

        Args:
            file: 上传的文件（.md 或 .docx）。

        Returns:
            ImportDocumentResponse 包含提取的标题/正文/标签/媒体。
        """
        raw_text, file_media = self._parse_upload(file)
        settings = get_settings()

        # LLM 深入提取
        extracted = self._extractor.extract(raw_text)

        # 合并文件级别的媒体信息
        merged_media: list[ImportedMedia] = []
        seen_names: set[str] = set()
        # 优先 LLM 提取的媒体
        for m in extracted.get("media", []):
            name = m.get("name", "")
            if name and name not in seen_names:
                merged_media.append(ImportedMedia(
                    index=m.get("index", 0),
                    name=name,
                    kind=m.get("kind", "image"),
                    description=m.get("description"),
                ))
                seen_names.add(name)
        # 补充文件解析出的媒体
        for m in file_media:
            name = m.get("name", "")
            if name and name not in seen_names:
                merged_media.append(ImportedMedia(
                    index=m.get("index", len(merged_media)),
                    name=name,
                    kind=m.get("kind", "image"),
                    description=m.get("description"),
                ))
                seen_names.add(name)

        # ---- 章节智能分析：段落/标题/子标题识别 ----
        body_text = extracted.get("body", raw_text)
        analysis = self._analyst.analyze(
            body=body_text,
            title=extracted.get("title", ""),
            tags=extracted.get("tags", []),
            content_type=extracted.get("content_type", "article"),
            allow_llm=settings.import_analysis_use_llm,
        )

        def _serialize_chapters(chapters: list[Any]) -> list[dict[str, Any]]:
            result: list[dict[str, Any]] = []
            for ch in chapters:
                item: dict[str, Any] = {
                    "level": ch.level,
                    "title": ch.title,
                    "content": ch.content,
                    "start_index": ch.start_index,
                    "word_count": ch.word_count,
                }
                if ch.sub_chapters:
                    item["sub_chapters"] = _serialize_chapters(ch.sub_chapters)
                result.append(item)
            return result

        return ImportDocumentResponse(
            title=extracted.get("title", ""),
            subtitle=analysis.subtitle or "",
            body=body_text,
            tags=extracted.get("tags", []),
            content_type=analysis.content_type or extracted.get("content_type", "article"),
            summary=analysis.summary or extracted.get("summary", ""),
            media=merged_media,
            chapters=_serialize_chapters(analysis.chapters or analysis.flat_chapters),
            raw_text=raw_text,
        )
